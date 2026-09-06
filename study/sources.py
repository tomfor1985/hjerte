import hashlib
import io
import re
import shutil
import subprocess
import zipfile
from pathlib import Path
from django.core.exceptions import ValidationError
from django.core.files.base import ContentFile
from django.db import transaction
from .models import Source, SourcePage, Chapter, Topic, Question

ALLOWED = {'.pdf', '.docx', '.pptx', '.txt', '.md'}


def normalize(text):
    return ' '.join(str(text).replace('\u00ad', '').split()).casefold()


def extract_pages(data, name):
    ext = Path(name).suffix.lower()
    if ext not in ALLOWED or len(data) > 50 * 1024 * 1024:
        raise ValidationError('Use a PDF, DOCX, PPTX, TXT or MD file up to 50 MB.')
    if ext in {'.docx', '.pptx'}:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 150 * 1024 * 1024:
                raise ValidationError('The document expands beyond the import limit.')
    if ext == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValidationError('Upload an unencrypted source document.')
        if len(reader.pages) > 600:
            raise ValidationError('Import documents with at most 600 pages.')
        if not shutil.which('pdftotext'):
            raise ValidationError('PDF import requires Poppler (pdftotext) on the server.')
        extracted=subprocess.run(['pdftotext','-layout','-','-'],input=data,capture_output=True,timeout=90,check=True)
        pages=extracted.stdout.decode('utf-8').split('\f')
        if pages and not pages[-1].strip(): pages.pop()
        if len(pages)!=len(reader.pages):
            raise ValidationError('Extracted PDF page count does not match the source.')
    elif ext == '.docx':
        from docx import Document
        doc = Document(io.BytesIO(data))
        chunks = []
        current = []
        for p in doc.paragraphs:
            if p.style.name.startswith('Heading') and current:
                chunks.append('\n'.join(current))
                current = []
            if p.text.strip():
                current.append(p.text)
        if current:
            chunks.append('\n'.join(current))
        for t in doc.tables:
            chunks.append('\n'.join(' | '.join(c.text for c in r.cells) for r in t.rows))
        pages = chunks
    elif ext == '.pptx':
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        pages = []
        for s in prs.slides:
            parts = [shape.text for shape in s.shapes if shape.has_text_frame]
            if s.has_notes_slide:
                parts.append(s.notes_slide.notes_text_frame.text)
            pages.append('\n'.join(parts))
    else:
        text = data.decode('utf-8-sig')
        pages = [text[i:i+6000] for i in range(0, len(text), 6000)]
    if not pages or len(''.join(pages).strip()) < 100:
        raise ValidationError('No usable text found. This document may need OCR before import.')
    if sum(map(len, pages)) > 3_000_000:
        raise ValidationError('Extracted text exceeds the import limit.')
    return [p.replace('\x00', '').replace('\u00ad', '') for p in pages]


def import_source(data, name, *, title='', kind='notes', year=None, doi='', url='', topic=None):
    digest = hashlib.sha256(data).hexdigest()
    existing = Source.objects.filter(sha256=digest).first()
    if existing:
        return existing, False
    pages = extract_pages(data, name)
    leading = '\n'.join(pages[:2])
    is_article=Path(name).suffix.lower()=='.pdf' and len(pages)>80
    if is_article and '10.1093/eurheartj/ehae178' in leading and '2024 esc guidelines for the management' in normalize(pages[0]):
        title, year, doi, kind = 'ESC elevated blood pressure and hypertension', 2024, '10.1093/eurheartj/ehae178', 'guideline'
    elif is_article and '10.1093/eurheartj/ehab484' in leading and '2021 esc guidelines on cardiovascular disease' in normalize(pages[0]):
        title, year, doi, kind = 'ESC cardiovascular disease prevention', 2021, '10.1093/eurheartj/ehab484', 'guideline'
    warnings = []
    empty = sum(not p.strip() for p in pages)
    if empty:
        warnings.append(f'{empty} pages have no extracted text; image-only material needs visual review.')
    if Path(name).suffix.lower() == '.docx':
        warnings.append('Locations are document sections, not printed page numbers.')
    source = Source(title=title or Path(name).stem, kind=kind, year=year, doi=doi,
                    url=url or (f'https://doi.org/{doi}' if doi else ''), sha256=digest,
                    original_name=Path(name).name, page_count=len(pages), extraction_warning=' '.join(warnings))
    try:
        with transaction.atomic():
            source.file.save(digest[:16]+Path(name).suffix.lower(), ContentFile(data), save=False)
            source.save()
            SourcePage.objects.bulk_create([SourcePage(source=source, number=i, text=t) for i,t in enumerate(pages, 1)])
            from .pdf_reading import prepare_source_reading
            prepare_source_reading(source)
            setup_known_chapters(source)
            if source.kind == 'guideline' and not source.chapters.exists() and topic:
                suggest_chapters(source, data, topic)
    except Exception:
        if source.file:
            source.file.delete(save=False)
        raise
    return source, True


def setup_known_chapters(source):
    if source.kind != 'guideline':
        return
    # Inclusive physical PDF page numbers, mapped from each source's contents.
    mapping = []
    if source.doi.endswith('ehae178'):
        mapping = [
            ('hypertension', 'Pathophysiology and clinical consequences', '3–4', 17, 18),
            ('hypertension', 'Measurement and diagnosis', '5', 18, 25),
            ('hypertension', 'Classification and cardiovascular risk', '6', 25, 31),
            ('hypertension', 'Work-up, resistant and secondary hypertension', '7', 32, 42),
            ('hypertension', 'Lifestyle and treatment strategies', '8', 43, 58),
            ('hypertension', 'Special populations and comorbidities', '9', 58, 72),
            ('hypertension', 'Acute blood pressure management and patient-centred care', '10–11', 72, 76),
        ]
    elif source.doi.endswith('ehab484'):
        mapping = [
            ('risk-assessment', 'Risk assessment and risk modifiers', '3', 9, 41),
            ('lifestyle', 'Communication, physical activity and nutrition', '4.1–4.3', 41, 46),
            ('psychosocial', 'Mental health and psychosocial risk', '4.4', 46, 47),
            ('smoking', 'Smoking cessation', '4.5', 47, 49),
            ('lipids', 'Lipids and lipoproteins', '4.6', 49, 55),
            ('diabetes', 'Diabetes and antithrombotic prevention', '4.8–4.9', 63, 66),
            ('public-health', 'Population prevention and public health', '5', 66, 69),
            ('rehabilitation', 'Disease-specific prevention and rehabilitation', '6', 69, 75),
        ]
    for slug,title,section,first,last in mapping:
        topic = Topic.objects.filter(slug=slug).first()
        if topic and last <= source.page_count:
            Chapter.objects.get_or_create(source=source, title=title,
                defaults={'topic':topic, 'section':section, 'first_page':first, 'last_page':last})


def suggest_chapters(source, data, topic):
    """Free local analysis: use PDF bookmarks, otherwise complete bounded ranges."""
    starts = {}
    if source.original_name.lower().endswith('.pdf'):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        for entry in reader.outline:
            if isinstance(entry, list):
                continue
            page = reader.get_destination_page_number(entry)
            if page is not None and 0 <= page < source.page_count:
                starts.setdefault(page + 1, str(entry.title)[:200])
    if 2 <= len(starts) <= 60:
        starts.setdefault(1, 'Introduction')
        ordered = sorted(starts)
        ranges = [(starts[p], p, ordered[i+1]-1 if i+1 < len(ordered) else source.page_count)
                  for i,p in enumerate(ordered)]
    else:
        ranges = [(f'Pages {p}–{min(p+7,source.page_count)}', p, min(p+7,source.page_count))
                  for p in range(1,source.page_count+1,8)]
        source.extraction_warning += ' No usable chapter bookmarks found. Suggested page groups cover the full document; you can rename or adjust them.'
        source.save(update_fields=['extraction_warning'])
    Chapter.objects.bulk_create([Chapter(source=source, topic=topic, title=title,
        first_page=first,last_page=last) for title,first,last in ranges])


@transaction.atomic
def retire_source(source,user):
    from django.utils import timezone
    from .models import CoverageSegment
    source.active=False
    source.save(update_fields=['active'])
    # Alternate formats represent the same edition. Withdrawing the main file
    # must also withdraw questions that retain an alternate file's provenance.
    count=sum(retire_source(copy,user) for copy in source.format_copies.filter(active=True))
    # Notes inventoried against this edition must not keep a green coverage
    # status after their authoritative evidence is withdrawn.
    for segment in CoverageSegment.objects.filter(status='mapped',audit__evidence_source_id=source.pk):
        segment.status='blocked'
        segment.audit={**segment.audit,'reason':'The supporting guideline was retired. Verify against its replacement.'}
        segment.save(update_fields=['status','audit'])
    page_ids=set(source.pages.values_list('id',flat=True))
    for q in Question.objects.exclude(status='retired').select_related('chapter'):
        provenance=q.verification.get('provenance',{})
        linked=(q.chapter.source_id==source.pk or provenance.get('source_id')==source.pk
                or source.pk in provenance.get('note_source_ids',[])
                or any(r.get('page_id') in page_ids for r in q.references))
        if linked:
            q.status='retired'
            q.verification={**q.verification,'retired_source_id':source.pk,
                'retired_at':timezone.now().isoformat(),'retired_by':user.username}
            q.save(update_fields=['status','verification','updated_at'])
            count+=1
    return count


def validate_references(refs, allowed_page_ids=None):
    if not isinstance(refs, list) or not refs:
        raise ValidationError('At least one verified guideline reference is required.')
    for ref in refs:
        if not isinstance(ref, dict):
            raise ValidationError('Reference must include a source page, section and supporting quotation.')
        if type(ref.get('page_id')) is not int:
            raise ValidationError('Reference page_id must be an integer.')
        page = SourcePage.objects.select_related('source').filter(id=ref.get('page_id')).first()
        if not page or not page.source.active or page.source.kind != 'guideline':
            raise ValidationError('Reference must point to an active guideline page in the library.')
        if allowed_page_ids is not None and page.pk not in allowed_page_ids:
            raise ValidationError('The model cited a page outside the supplied source context.')
        quote = normalize(ref.get('quote', ''))
        content,minimum=page.text,30
        if ref.get('reading_sha256'):
            from .pdf_reading import reading_for
            reading=reading_for(page)
            if not reading or reading.text_sha256!=ref['reading_sha256']:
                raise ValidationError('The reference reading view is missing or no longer matches its source.')
            if '\ufffd' in quote:
                raise ValidationError('The source passage contains unreadable PDF glyphs. Inspect the original or OCR before using it as evidence.')
            if not any(p['start']==ref.get('passage_start') and p['end']==ref.get('passage_end') and
                       quote==normalize(reading.text[p['start']:p['end']]) for p in reading.passages):
                raise ValidationError('The reference is not an exact saved PDF passage.')
            content,minimum=reading.text,1
        if len(quote) < minimum or len(quote) > 900 or quote not in normalize(content):
            raise ValidationError('Supporting quotation does not match the cited source page.')
        if not str(ref.get('section', '')).strip():
            raise ValidationError('A guideline section must be specified.')


def reference_snapshot(refs):
    result = []
    for ref in refs:
        p = SourcePage.objects.select_related('source').get(pk=ref['page_id'])
        result.append({**ref, 'source_id':p.source_id, 'title':p.source.title, 'year':p.source.year,
                       'doi':p.source.doi, 'url':p.source.url, 'page':p.number, 'sha256':p.source.sha256})
    return result
